from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
import re
from typing import Literal

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TemplateFieldType = Literal["text", "textarea", "dropdown", "multi_select", "tone_weights", "color_pairs"]


@dataclass(frozen=True)
class BrandSpaceTemplateField:
    key: str
    section: str
    label: str
    field_type: TemplateFieldType
    options: tuple[str, ...] = ()
    value_map: tuple[tuple[str, str], ...] = ()


# This explicit registry controls the document surface. It includes every active,
# manually fillable Brand Space creation field from the approved inventory.
# File uploads, generated values, and hidden/legacy form state stay excluded.
BRAND_SPACE_TEMPLATE_FIELDS: tuple[BrandSpaceTemplateField, ...] = (
    # Core Brand Signals
    BrandSpaceTemplateField("core.name", "Core Brand Signals", "Brand Name", "text"),
    BrandSpaceTemplateField("core.tagline", "Core Brand Signals", "Tagline", "text"),
    BrandSpaceTemplateField("core.description", "Core Brand Signals", "Brand Description", "textarea"),
    BrandSpaceTemplateField("core.industryCategory", "Core Brand Signals", "Industry Category", "dropdown", (
        "Technology / SaaS", "Financial Services", "Healthcare", "Retail/E-commerce", "FMCG", "Education",
        "Real Estate", "Automotive", "Hospitality", "Media & Entertainment", "Manufacturing", "Energy",
        "Telecommunications", "Professional Services", "Government / Public Sector", "Nonprofit / NGO", "other",
    )),
    BrandSpaceTemplateField("core.differentiators", "Core Brand Signals", "Key Differentiators", "textarea"),

    # Brand Foundations
    BrandSpaceTemplateField("additional.businessModels", "Brand Foundations: Business Model", "Business Model", "multi_select", ("B2B", "B2C", "B2B2C", "Other")),
    BrandSpaceTemplateField("additional.businessModelDetails.b2b", "Brand Foundations: Business Model", "B2B Detail", "text"),
    BrandSpaceTemplateField("additional.businessModelDetails.b2c", "Brand Foundations: Business Model", "B2C Detail", "text"),
    BrandSpaceTemplateField("additional.businessModelDetails.b2b2c", "Brand Foundations: Business Model", "B2B2C Detail", "text"),
    BrandSpaceTemplateField("additional.businessModelDetails.other", "Brand Foundations: Business Model", "Other Detail", "text"),
    BrandSpaceTemplateField("additional.businessModelOther", "Brand Foundations: Business Model", "Other Business Model", "text"),
    BrandSpaceTemplateField("additional.routesToMarket", "Brand Foundations: Route to Market", "Route to Market", "multi_select", (
        "D2C", "Retail", "Marketplace", "Distributor/Dealer", "Partner-led", "Direct Sales", "Other",
    )),
    BrandSpaceTemplateField("additional.routeToMarketDetails.d2c", "Brand Foundations: Route to Market", "D2C Detail", "text"),
    BrandSpaceTemplateField("additional.routeToMarketDetails.retail", "Brand Foundations: Route to Market", "Retail Detail", "text"),
    BrandSpaceTemplateField("additional.routeToMarketDetails.marketplace", "Brand Foundations: Route to Market", "Marketplace Detail", "text"),
    BrandSpaceTemplateField("additional.routeToMarketDetails.distributor_dealer", "Brand Foundations: Route to Market", "Distributor/Dealer Detail", "text"),
    BrandSpaceTemplateField("additional.routeToMarketDetails.partner_led", "Brand Foundations: Route to Market", "Partner-led Detail", "text"),
    BrandSpaceTemplateField("additional.routeToMarketDetails.direct_sales", "Brand Foundations: Route to Market", "Direct Sales Detail", "text"),
    BrandSpaceTemplateField("additional.routeToMarketDetails.other", "Brand Foundations: Route to Market", "Other Detail", "text"),
    BrandSpaceTemplateField("additional.brandMission", "Brand Foundations: Purpose & Positioning", "Brand Mission", "text"),
    BrandSpaceTemplateField("additional.brandVision", "Brand Foundations: Purpose & Positioning", "Brand Vision", "text"),
    BrandSpaceTemplateField("additional.brandPromise", "Brand Foundations: Purpose & Positioning", "Brand Promise", "text"),
    BrandSpaceTemplateField("additional.marketPositioning", "Brand Foundations: Purpose & Positioning", "Market Positioning", "text"),
    BrandSpaceTemplateField("additional.roleOfDigitalPlatforms", "Brand Foundations: Purpose & Positioning", "Role of Digital Platforms", "text"),
    BrandSpaceTemplateField("additional.socialMediaChallenges", "Brand Foundations: Purpose & Positioning", "Social Media Challenges", "text"),
    BrandSpaceTemplateField("additional.businessProblemOrOpportunity", "Brand Foundations: Strategic Block", "Business Problem or Opportunity", "text"),
    BrandSpaceTemplateField("additional.perceptionChallenge", "Brand Foundations: Strategic Block", "Perception Challenge", "text"),
    BrandSpaceTemplateField("additional.humanInsight", "Brand Foundations: Strategic Block", "Human Insight", "text"),
    BrandSpaceTemplateField("additional.brandAdvantage", "Brand Foundations: Strategic Block", "Brand Advantage", "text"),
    BrandSpaceTemplateField("additional.strategy", "Brand Foundations: Strategic Block", "Strategy", "textarea"),
    BrandSpaceTemplateField("additional.brandArchetype", "Brand Foundations: Industry & Context", "Brand Archetype", "dropdown", (
        "Hero", "Innovator", "Caregiver", "Explorer", "Creator", "Sage", "Rebel", "Entertainer", "Everyman", "Ruler", "Lover", "Magician",
    )),
    BrandSpaceTemplateField("additional.complianceLevel", "Brand Foundations: Industry & Context", "Compliance Level", "dropdown", ("Low", "Medium", "High")),
    BrandSpaceTemplateField("additional.competitorBrands.0.name", "Brand Foundations: Competitor 1", "Competitor Brand Name", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.0.websiteUrl", "Brand Foundations: Competitor 1", "Website URL", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.0.linkedin", "Brand Foundations: Competitor 1", "LinkedIn", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.0.instagram", "Brand Foundations: Competitor 1", "Instagram", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.0.x", "Brand Foundations: Competitor 1", "X", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.1.name", "Brand Foundations: Competitor 2", "Competitor Brand Name", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.1.websiteUrl", "Brand Foundations: Competitor 2", "Website URL", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.1.linkedin", "Brand Foundations: Competitor 2", "LinkedIn", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.1.instagram", "Brand Foundations: Competitor 2", "Instagram", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.1.x", "Brand Foundations: Competitor 2", "X", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.2.name", "Brand Foundations: Competitor 3", "Competitor Brand Name", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.2.websiteUrl", "Brand Foundations: Competitor 3", "Website URL", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.2.linkedin", "Brand Foundations: Competitor 3", "LinkedIn", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.2.instagram", "Brand Foundations: Competitor 3", "Instagram", "text"),
    BrandSpaceTemplateField("additional.competitorBrands.2.x", "Brand Foundations: Competitor 3", "X", "text"),

    # Brand Voice & Emotion
    BrandSpaceTemplateField("voiceTone.coreToneAttributes", "Brand Voice & Emotion", "Core Tone Attributes", "multi_select", (
        "Professional / Formal", "Bold", "Premium", "Playful", "Authoritative", "Empathetic", "Inspirational", "Trust Worthy", "Polite", "Witty",
    )),
    BrandSpaceTemplateField("voiceTone.coreToneAttributeWeights", "Brand Voice & Emotion", "Tone Weights", "tone_weights", (
        "Professional / Formal", "Bold", "Premium", "Playful", "Authoritative", "Empathetic", "Inspirational", "Trust Worthy", "Polite", "Witty",
    )),
    BrandSpaceTemplateField("voiceTone.primaryEmotion", "Brand Voice & Emotion", "Primary Emotion", "text"),
    BrandSpaceTemplateField("voiceTone.secondaryEmotion", "Brand Voice & Emotion", "Secondary Emotion", "text"),
    BrandSpaceTemplateField("voiceTone.avoidedEmotion", "Brand Voice & Emotion", "Avoided Emotion", "text"),
    BrandSpaceTemplateField("voiceTone.contentComplexity", "Brand Voice & Emotion", "Content Complexity", "dropdown", ("Basic", "Expert")),
    BrandSpaceTemplateField("voiceTone.sentenceLength", "Brand Voice & Emotion", "Sentence Length", "dropdown", ("Short", "Medium", "Long", "Mixed")),

    # Audience Persona Mapping
    BrandSpaceTemplateField("targetAudience.selectedAudiences", "Audience Persona Mapping: Audience", "Audience Names", "textarea"),
    BrandSpaceTemplateField("targetAudience.goals", "Audience Persona Mapping: Psychographics", "Goals", "textarea"),
    BrandSpaceTemplateField("targetAudience.motivations", "Audience Persona Mapping: Psychographics", "Motivations", "textarea"),
    BrandSpaceTemplateField("targetAudience.fearsAndPainPoints", "Audience Persona Mapping: Psychographics", "Fears and Pain Points", "textarea"),
    BrandSpaceTemplateField("targetAudience.objections", "Audience Persona Mapping: Psychographics", "Objections", "textarea"),
    BrandSpaceTemplateField("targetAudience.contentConsumptionBehavior", "Audience Persona Mapping: Psychographics", "Content Consumption Behavior", "textarea"),
    BrandSpaceTemplateField("targetAudience.audienceType", "Audience Persona Mapping: Demographics", "Audience Type", "dropdown", ("Consumer", "Professional")),
    BrandSpaceTemplateField("targetAudience.location", "Audience Persona Mapping: Demographics", "Location/Region", "dropdown", ("Local", "Global")),
    BrandSpaceTemplateField("targetAudience.locationDetail", "Audience Persona Mapping: Demographics", "Location Detail", "text"),
    BrandSpaceTemplateField("targetAudience.educationLevel", "Audience Persona Mapping: Demographics", "Education Level", "dropdown", (
        "High School or Below", "College / Diploma Educated", "University Graduate / Postgraduate Educated", "Highly Educated / Academic",
    )),
    BrandSpaceTemplateField("targetAudience.employmentStatus", "Audience Persona Mapping: Demographics", "Employment Status", "dropdown", (
        "Student", "Early Career Professional", "Mid Career Professional", "Senior Professional", "Executive / Leadership", "Entrepreneur / Business Owner", "Freelancer / Independent Worker", "Homemaker", "Retired",
    )),
    BrandSpaceTemplateField("targetAudience.professionalBackground", "Audience Persona Mapping: Demographics", "Professional Background", "dropdown", (
        "Technology", "Business", "Finance", "Healthcare", "Education", "Creative", "Sales", "Operations", "Legal", "Entrepreneurship", "Skilled Trades", "Government", "Student", "General Audience",
    )),
    BrandSpaceTemplateField("targetAudience.householdSize", "Audience Persona Mapping: Demographics", "Household Size", "dropdown", (
        "Single Person Household", "Couple Household", "Small Family (3-4 Members)", "Large Family (5+ Members)", "Shared / Multi Generational Household",
    )),
    BrandSpaceTemplateField("targetAudience.languagePreference", "Audience Persona Mapping: Demographics", "Language Preference", "dropdown", (
        "Local Language", "Local Language + English", "English Preferred", "Multilingual Audience",
    )),
    BrandSpaceTemplateField("targetAudience.incomeLevel", "Audience Persona Mapping: Demographics", "Income Level", "dropdown", (
        "Low Income", "Lower Middle Income", "Middle Income", "Upper Middle Income", "High Income",
    )),
    BrandSpaceTemplateField("targetAudience.familyStatusOrLifeStage", "Audience Persona Mapping: Demographics", "Family Status or Life Stage", "text"),
    BrandSpaceTemplateField("targetAudience.socioEconomicSegment", "Audience Persona Mapping: Demographics", "Socio-economic Segment", "text"),
    BrandSpaceTemplateField("targetAudience.digitalAccess", "Audience Persona Mapping: Demographics", "Digital Access", "dropdown", (
        "Mobile Only", "Mobile First", "Mobile and Desktop", "Multi Device Power User",
    )),

    # Do's & Don'ts
    BrandSpaceTemplateField("brandRules.selectedRules", "Do's & Don'ts", "Selected Rules", "multi_select", (
        "Avoid emojis", "Avoid hype claims", "Avoid slang", "Avoid competitor comparison",
    )),
    BrandSpaceTemplateField("brandRules.positiveWordBank", "Do's & Don'ts", "Positive Word Bank", "textarea"),
    BrandSpaceTemplateField("brandRules.replaceableWords", "Do's & Don'ts", "Replaceable Words", "textarea"),
    BrandSpaceTemplateField("brandRules.negativeWordBank", "Do's & Don'ts", "Negative Word Bank", "textarea"),
    BrandSpaceTemplateField("brandRules.whatToDo", "Do's & Don'ts", "What To Do", "textarea"),
    BrandSpaceTemplateField("brandRules.whatNotToDo", "Do's & Don'ts", "What NOT To Do", "textarea"),
    BrandSpaceTemplateField("brandRules.restrictedTopics", "Do's & Don'ts", "Restricted Topics", "textarea"),
    BrandSpaceTemplateField("brandRules.restrictedClaims", "Do's & Don'ts", "Restricted Claims", "textarea"),
    BrandSpaceTemplateField("brandRules.permittedClaims", "Do's & Don'ts", "Permitted Claims", "textarea"),
    BrandSpaceTemplateField("brandRules.blockedWordsPhrases", "Do's & Don'ts", "Blocked Words / Phrases", "textarea"),

    # Prompt Intelligence Setup
    BrandSpaceTemplateField("promptIntelligence.preferredPlatforms", "Prompt Intelligence Setup", "Preferred Platforms", "multi_select", (
        "LinkedIn", "Instagram", "X (Twitter)", "YouTube", "Facebook", "TikTok", "Pinterest", "Threads",
    )),
    BrandSpaceTemplateField("promptIntelligence.contentFormats", "Prompt Intelligence Setup", "Preferred Content Formats", "multi_select", (
        "Short-form post", "Long-form article", "Carousel", "Reel / Short video", "Story", "Newsletter", "Thread", "Infographic caption",
    )),
    BrandSpaceTemplateField("promptIntelligence.contentTone", "Prompt Intelligence Setup", "Content Tone Override", "text"),
    BrandSpaceTemplateField("promptIntelligence.platformRules", "Prompt Intelligence Setup", "Platform-Specific Rules", "textarea"),
    BrandSpaceTemplateField("promptIntelligence.contextualHints", "Prompt Intelligence Setup", "Contextual Hints", "textarea"),
    BrandSpaceTemplateField("promptIntelligence.instructionOverrides", "Prompt Intelligence Setup", "Instruction Overrides (Global)", "textarea"),
    BrandSpaceTemplateField("promptIntelligence.avoidedFormats", "Prompt Intelligence Setup", "Formats to Avoid", "text"),

    # Content Objectives
    BrandSpaceTemplateField("objectives.primaryObjective", "Content Objectives", "Primary Objective", "dropdown", (
        "Brand Awareness", "Lead Generation", "Sales Conversion", "Community Building", "Product Launch", "Thought Leadership", "Customer Retention", "Employer Branding",
    ), (
        ("Brand Awareness", "brand_awareness"), ("Lead Generation", "lead_generation"), ("Sales Conversion", "sales_conversion"), ("Community Building", "community_building"),
        ("Product Launch", "product_launch"), ("Thought Leadership", "thought_leadership"), ("Customer Retention", "customer_retention"), ("Employer Branding", "employer_branding"),
    )),
    BrandSpaceTemplateField("objectives.contentGoal", "Content Objectives", "Content Goal", "dropdown", (
        "Educate", "Inspire", "Entertain", "Convert", "Inform", "Engage", "Nurture",
    ), (("Educate", "educate"), ("Inspire", "inspire"), ("Entertain", "entertain"), ("Convert", "convert"), ("Inform", "inform"), ("Engage", "engage"), ("Nurture", "nurture"))),
    BrandSpaceTemplateField("objectives.campaignTheme", "Content Objectives", "Campaign Theme", "text"),
    BrandSpaceTemplateField("objectives.businessOutcome", "Content Objectives", "Business Outcome", "textarea"),
    BrandSpaceTemplateField("objectives.callToAction", "Content Objectives", "Call to Action", "text"),
    BrandSpaceTemplateField("objectives.targetConversionAction", "Content Objectives", "Target Conversion Action", "text"),
    BrandSpaceTemplateField("objectives.contentFrequency", "Content Objectives", "Content Frequency", "dropdown", (
        "Daily", "3x per week", "Weekly", "Bi-weekly", "Monthly", "Campaign-based",
    ), (("Daily", "daily"), ("3x per week", "3x_week"), ("Weekly", "weekly"), ("Bi-weekly", "biweekly"), ("Monthly", "monthly"), ("Campaign-based", "campaign_based"))),
    BrandSpaceTemplateField("objectives.successMetric", "Content Objectives", "Primary Success Metric", "dropdown", (
        "Reach / Impressions", "Engagement Rate", "Clicks / CTR", "Conversions", "Follower Growth", "Share of Voice", "Revenue Impact",
    ), (("Reach / Impressions", "reach"), ("Engagement Rate", "engagement_rate"), ("Clicks / CTR", "clicks"), ("Conversions", "conversions"), ("Follower Growth", "followers"), ("Share of Voice", "share_of_voice"), ("Revenue Impact", "revenue"))),

    # Visual Identity
    BrandSpaceTemplateField("visualIdentity.brandMood", "Visual Identity", "Brand Mood", "textarea"),
    BrandSpaceTemplateField("visualIdentity.visualStyle", "Visual Identity", "Visual Style", "textarea"),
    BrandSpaceTemplateField("visualIdentity.logoPlacements", "Visual Identity", "Logo Placement", "dropdown", (
        "Top - Right", "Top - Left", "Top - Center", "Bottom - Right", "Bottom - Left", "Bottom - Center", "Center",
    )),
    BrandSpaceTemplateField("visualIdentity.primaryColor", "Visual Identity", "Primary Color", "text"),
    BrandSpaceTemplateField("visualIdentity.secondaryColor", "Visual Identity", "Secondary Color", "text"),
    BrandSpaceTemplateField("visualIdentity.additionalColors", "Visual Identity", "Additional Colors", "color_pairs"),
    BrandSpaceTemplateField("visualIdentity.typography", "Visual Identity", "Font", "text"),
)

_FIELD_BY_KEY = {field.key: field for field in BRAND_SPACE_TEMPLATE_FIELDS}
_MARKER_PATTERN = re.compile(r"\[\[violyt:([A-Za-z0-9_.-]+)\]\]", re.IGNORECASE)


def _set_font(run, *, size: int, bold: bool = False, color: str = "1F1F1F") -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def _set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_width = cell._tc.tcPr.tcW
    tc_width.set(qn("w:w"), str(round(width_inches * 1440)))
    tc_width.set(qn("w:type"), "dxa")


class BrandSpaceTemplateService:
    filename = "brand-space-template.docx"

    def build_document(self) -> bytes:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        normal.font.size = Pt(10)
        normal.paragraph_format.space_after = Pt(5)

        title = document.add_paragraph()
        title.paragraph_format.space_after = Pt(3)
        _set_font(title.add_run("Brand Space Template"), size=22, bold=True, color="3C2F8F")
        instructions = document.add_paragraph()
        instructions.paragraph_format.space_after = Pt(14)
        _set_font(
            instructions.add_run(
                "Complete only the fields you need. For multi-select fields, separate choices with commas. "
                "For tone weights, use Tone: 0-100. For additional colors, use Color name: #RRGGBB, one per line. "
                "Keep the field labels unchanged so Violyt can import your answers."
            ),
            size=10,
            color="5F6368",
        )

        current_section = None
        table = None
        for field in BRAND_SPACE_TEMPLATE_FIELDS:
            if field.section != current_section:
                current_section = field.section
                heading = document.add_paragraph()
                heading.paragraph_format.space_before = Pt(12)
                heading.paragraph_format.space_after = Pt(5)
                _set_font(heading.add_run(current_section), size=14, bold=True, color="3C2F8F")
                table = document.add_table(rows=0, cols=2)
                table.autofit = False
                table.style = "Table Grid"

            assert table is not None
            row = table.add_row()
            label_cell, value_cell = row.cells
            _set_cell_width(label_cell, 2.25)
            _set_cell_width(value_cell, 4.75)
            label_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            label = label_cell.paragraphs[0]
            label.paragraph_format.space_after = Pt(2)
            _set_font(label.add_run(field.label), size=10, bold=True)
            if field.options:
                option_paragraph = label_cell.add_paragraph()
                option_paragraph.paragraph_format.space_after = Pt(0)
                _set_font(option_paragraph.add_run("Choices: " + "; ".join(field.options)), size=8, color="5F6368")
            marker = label_cell.add_paragraph()
            marker.paragraph_format.space_after = Pt(0)
            marker_run = marker.add_run(f"[[violyt:{field.key}]]")
            marker_run.font.hidden = True

            value_paragraph = value_cell.paragraphs[0]
            value_paragraph.paragraph_format.space_after = Pt(0)
            value_paragraph.paragraph_format.line_spacing = 1.15
            if field.field_type == "tone_weights":
                placeholder = "Example: Bold: 70, Trust Worthy: 40"
            elif field.field_type == "color_pairs":
                placeholder = "Example: Accent: #112233 (one color per line)"
            elif field.field_type in {"text", "dropdown"}:
                placeholder = "Type your response"
            else:
                placeholder = "Type one or more responses"
            _set_font(value_paragraph.add_run(placeholder), size=10, color="A0A0A0")
            if field.field_type == "textarea":
                value_paragraph.add_run().add_break(WD_BREAK.LINE)
                value_paragraph.add_run().add_break(WD_BREAK.LINE)

        output = BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _normalize_value(
        field: BrandSpaceTemplateField,
        value: str,
    ) -> str | list[str] | dict[str, int] | list[dict[str, str]] | None:
        cleaned = value.replace("\u00a0", " ").strip()
        if not cleaned or cleaned.lower() in {
            "type your response",
            "type one or more responses",
            "example: bold: 70, trust worthy: 40",
            "example: accent: #112233 (one color per line)",
        }:
            return None
        if field.field_type == "textarea" or field.field_type == "text":
            return cleaned

        if field.field_type == "dropdown":
            matched = next((option for option in field.options if option.casefold() == cleaned.casefold()), None)
            if matched:
                return dict(field.value_map).get(matched, matched)
            return next(
                (stored_value for _, stored_value in field.value_map if stored_value.casefold() == cleaned.casefold()),
                None,
            )

        if field.field_type == "tone_weights":
            weights: dict[str, int] = {}
            for item in re.split(r"[,;\n]+", cleaned):
                name, separator, raw_weight = item.partition(":")
                if not separator:
                    continue
                option = next((choice for choice in field.options if choice.casefold() == name.strip().casefold()), None)
                numeric_weight = raw_weight.strip().removesuffix("%").strip()
                if not option or not numeric_weight.isdigit():
                    continue
                weight = int(numeric_weight)
                if 0 <= weight <= 100:
                    weights[option] = weight
            return weights or None

        if field.field_type == "color_pairs":
            colors: list[dict[str, str]] = []
            for item in re.split(r"[;\n]+", cleaned):
                name, separator, raw_hex = item.partition(":")
                hex_value = raw_hex.strip()
                if not separator or not name.strip() or not re.fullmatch(r"#[0-9a-fA-F]{6}", hex_value):
                    continue
                colors.append({"name": name.strip(), "hex": hex_value.upper()})
            return colors or None

        requested = [item.strip() for item in re.split(r"[,;\n]+", cleaned) if item.strip()]
        normalized = [
            option
            for option in field.options
            if any(candidate.casefold() == option.casefold() for candidate in requested)
        ]
        return normalized or None

    def parse_document(self, content: bytes) -> list[dict[str, object]]:
        try:
            document = Document(BytesIO(content))
        except Exception as exc:  # python-docx uses several package-specific errors for invalid uploads.
            raise ValueError("The uploaded file is not a readable .docx Brand Space template.") from exc

        imported: list[dict[str, object]] = []
        seen: set[str] = set()
        for table in document.tables:
            for row in table.rows:
                if len(row.cells) < 2:
                    continue
                match = _MARKER_PATTERN.search(row.cells[0].text)
                if not match:
                    continue
                key = match.group(1)
                field = _FIELD_BY_KEY.get(key)
                if not field or key in seen:
                    continue
                value = self._normalize_value(field, row.cells[1].text)
                if value is None:
                    continue
                seen.add(key)
                imported.append({"key": key, "value": value, "field_type": field.field_type})

        if not imported:
            raise ValueError("No completed Brand Space template fields were found in this document.")
        return imported
