from io import BytesIO

import pytest
from docx import Document

from app.services.brand_space_template import BRAND_SPACE_TEMPLATE_FIELDS, BrandSpaceTemplateService


def _set_template_value(document: Document, key: str, value: str) -> None:
    marker = f"[[violyt:{key}]]"
    for table in document.tables:
        for row in table.rows:
            if marker in row.cells[0].text:
                row.cells[1].text = value
                return
    raise AssertionError(f"Template field {key} was not found")


def test_template_document_contains_only_registered_fields() -> None:
    service = BrandSpaceTemplateService()
    document = Document(BytesIO(service.build_document()))
    markers = [
        paragraph.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if "[[violyt:" in paragraph.text
    ]

    assert len(markers) == len(BRAND_SPACE_TEMPLATE_FIELDS)
    assert "[[violyt:core.name]]" in markers
    assert "[[violyt:additional.competitorBrands.2.x]]" in markers
    assert "[[violyt:brandRules.permittedClaims]]" in markers
    assert "[[violyt:targetAudience.locationDetail]]" in markers
    assert "[[violyt:voiceTone.perspective]]" not in markers
    assert "[[violyt:brandKnowledge.templateFiles]]" not in markers


def test_template_parser_reads_only_filled_registered_fields() -> None:
    service = BrandSpaceTemplateService()
    document = Document(BytesIO(service.build_document()))
    _set_template_value(document, "core.name", "Northstar Labs")
    _set_template_value(document, "voiceTone.contentComplexity", "expert")
    _set_template_value(document, "voiceTone.coreToneAttributeWeights", "Bold: 70; Witty: 30; Invalid: 10")
    _set_template_value(document, "targetAudience.selectedAudiences", "Founders, Investors, Invalid Choice")
    _set_template_value(document, "targetAudience.location", "Global")
    _set_template_value(document, "targetAudience.locationDetail", "India")
    _set_template_value(document, "brandRules.selectedRules", "Avoid hype claims; Avoid slang")
    _set_template_value(document, "brandRules.permittedClaims", "Verified product benefits")
    _set_template_value(document, "objectives.primaryObjective", "Lead Generation")
    _set_template_value(document, "objectives.contentFrequency", "3x per week")
    _set_template_value(document, "additional.competitorBrands.1.websiteUrl", "https://example.com")
    _set_template_value(document, "visualIdentity.additionalColors", "Accent: #112233\nHighlight: #AABBCC")

    output = BytesIO()
    document.save(output)
    fields = {item["key"]: item["value"] for item in service.parse_document(output.getvalue())}

    assert fields == {
        "core.name": "Northstar Labs",
        "voiceTone.contentComplexity": "Expert",
        "voiceTone.coreToneAttributeWeights": {"Bold": 70, "Witty": 30},
        "targetAudience.selectedAudiences": ["Founders", "Investors"],
        "targetAudience.location": "Global",
        "targetAudience.locationDetail": "India",
        "brandRules.selectedRules": ["Avoid hype claims", "Avoid slang"],
        "brandRules.permittedClaims": "Verified product benefits",
        "objectives.primaryObjective": "lead_generation",
        "objectives.contentFrequency": "3x_week",
        "additional.competitorBrands.1.websiteUrl": "https://example.com",
        "visualIdentity.additionalColors": [
            {"name": "Accent", "hex": "#112233"},
            {"name": "Highlight", "hex": "#AABBCC"},
        ],
    }


def test_template_round_trips_brand_foundations_and_permitted_claims() -> None:
    service = BrandSpaceTemplateService()
    document = Document(BytesIO(service.build_document()))
    values = {
        "additional.businessModels": "B2B, B2C, B2B2C, Other",
        "additional.businessModelDetails.b2b": "Enterprise buyers",
        "additional.businessModelDetails.b2c": "Individual consumers",
        "additional.businessModelDetails.b2b2c": "Partner-led consumer offering",
        "additional.businessModelDetails.other": "Government",
        "additional.routesToMarket": "D2C, Retail, Marketplace, Distributor/Dealer, Partner-led, Direct Sales, Other",
        "additional.routeToMarketDetails.d2c": "Brand website",
        "additional.routeToMarketDetails.retail": "Retail stores",
        "additional.routeToMarketDetails.marketplace": "Online marketplaces",
        "additional.routeToMarketDetails.distributor_dealer": "Dealer network",
        "additional.routeToMarketDetails.partner_led": "Reseller partners",
        "additional.routeToMarketDetails.direct_sales": "Direct sales team",
        "additional.routeToMarketDetails.other": "Industry events",
        "brandRules.permittedClaims": "Claims supported by verified evidence",
    }
    for key, value in values.items():
        _set_template_value(document, key, value)

    output = BytesIO()
    document.save(output)
    fields = {item["key"]: item["value"] for item in service.parse_document(output.getvalue())}

    assert fields == {
        **values,
        "additional.businessModels": ["B2B", "B2C", "B2B2C", "Other"],
        "additional.routesToMarket": [
            "D2C",
            "Retail",
            "Marketplace",
            "Distributor/Dealer",
            "Partner-led",
            "Direct Sales",
            "Other",
        ],
    }


def test_template_parser_rejects_documents_without_completed_template_fields() -> None:
    document = Document()
    document.add_paragraph("An unrelated document")
    output = BytesIO()
    document.save(output)

    with pytest.raises(ValueError, match="No completed Brand Space template fields"):
        BrandSpaceTemplateService().parse_document(output.getvalue())
